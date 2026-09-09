"""
Settle Old Unpaid Invoices Against Per-Salesman "Old Balance" Ledgers
--------------------------------------------------------------------
For every aged unpaid invoice in ../data_change/old_bills.xlsx, posts a
Zoho Books *manual journal* that moves the invoice's open balance off the
customer's A/R onto a per-salesman Other-Current-Asset sub-account of a new
parent account "OLD BALANCE RECEIVABLES", then applies that journal's
receivables credit to the specific invoice.

Per invoice:
  1. POST /journals
       credit  Accounts Receivable   <balance>   (customer_id = the customer)
       debit   <SALESMAN> OLD BALANCE <balance>
     (org auto-publishes; falls back to POST /journals/{id}/status/publish)
  2. GET  /journals/{id}/credits            -> journal_line_id of the AR credit
  3. POST /journals/{id}/credits/{journal_line_id}/invoices
       {"invoices":[{"invoice_id","amount_applied"}]}
     -> invoice flips to Paid.

Effect per invoice:  Dr <SALESMAN> OLD BALANCE  /  Cr Accounts Receivable
Original revenue / VAT are untouched (debt transfer, not a write-off).

NOTE: Zoho *Payment Received* was tried first and rejected with
"Involved account types are not applicable" (deposit-to must be bank/cash/
undeposited-funds, not an Other Current Asset). The journal route is the
working one. `zohodocs/journals.yml` is stale and omits the
/journals/{id}/credits endpoints - they exist (zoho.com/books/api/v3/journals/).

Modeled on writeoff_invoices.py (same TokenManager / RateLimiter / QuotaTracker /
crash-safe jsonl logging / --undo conventions).

Every run that touches the API writes a row-by-row CSV under settle_runs/ (flushed
after each call, so a crash still leaves a full record):
  accounts_log_<ts>.csv  - each account created / reused
  settle_log_<ts>.csv    - one row per invoice journal attempt (+ .jsonl + .json)
  undo_log_<ts>.csv      - one row per journal deletion / account deactivation

Usage:
    pip install requests python-dotenv openpyxl

    python settle_old_bills_to_salesman.py --build
    python settle_old_bills_to_salesman.py --dry-run --input candidates_<ts>.json
    python settle_old_bills_to_salesman.py --input candidates_<ts>.json --limit 5
    python settle_old_bills_to_salesman.py --input candidates_<ts>.json --ledger-map ledger_map_<ts>.json
    python settle_old_bills_to_salesman.py --input candidates_<ts>.json --yes
    python settle_old_bills_to_salesman.py --undo settle_log_<ts>.json
    python settle_old_bills_to_salesman.py --undo settle_log_<ts>.json --undo-accounts --ledger-map ledger_map_<ts>.json
"""

import argparse
import csv
import json
import os
import re
import threading
import time
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed

import requests
from dotenv import load_dotenv

load_dotenv(os.path.join(os.path.dirname(__file__), "..", ".env_zoho"))

# --------------------------------------------------------------------------- #
# Config
# --------------------------------------------------------------------------- #
INPUT_XLSX = os.path.join(os.path.dirname(__file__), "..", "data_change", "old_bills.xlsx")
XLSX_SHEET = "Sheeet 1"

PARENT_ACCOUNT = "OLD BALANCE RECEIVABLES"
ACCOUNT_TYPE = "other_current_asset"
# This org makes account_code mandatory. Parent sits in the 1.8.x "Receivables"
# group; the 9 sub-accounts get 1.8.5100..1.8.5900 assigned alphabetically.
PARENT_ACCOUNT_CODE = "1.8.5000"
SUBACCOUNT_CODE_BASE = 5100      # 1.8.51xx .. stepped by 100 per ledger, sorted by name
JOURNAL_DATE = "2026-09-10"     # old periods are locked; all journals dated here
REF_PREFIX = "OLDBAL-JC"        # reference_number = f"{REF_PREFIX}-{invoice_number}"
AR_ACCOUNT_ID = None            # resolved at runtime from the chart of accounts

EXPECTED_TOTAL = 314104.25       # sanity check on the sheet
TOTAL_TOLERANCE = 5.0

RUNS_DIR = os.path.join(os.path.dirname(__file__), "settle_runs")
os.makedirs(RUNS_DIR, exist_ok=True)

# xlsx column indexes (0-based)
# Invoice#, Customer, Salesman, Current Salesman, Invoice Date, Due Date,
# Age (Days), Total Amount (AED), BALANCE TO SALESMAN, Balance (AED), Status,
# Customer ID, Invoice ID
C_INV, C_CUST, C_CURSALES, C_BAL, C_CUSTID, C_INVID = 0, 1, 3, 9, 11, 12
C_SALES, C_INVDATE, C_DUEDATE, C_AGE, C_TOTAL = 2, 4, 5, 6, 7


def _require_env(key: str) -> str:
    val = os.getenv(key)
    if not val:
        raise RuntimeError(f"Missing required env var: {key}  (set it in .env_zoho)")
    return val


CLIENT_ID = _require_env("ZOHO_CLIENT_ID")
CLIENT_SECRET = _require_env("ZOHO_CLIENT_SECRET")
REFRESH_TOKEN = _require_env("ZOHO_REFRESH_TOKEN")
ORG_ID = _require_env("ZOHO_ORG_ID")
TOKEN_URL = "https://accounts.zoho.com/oauth/v2/token"
API_BASE = "https://www.zohoapis.com/books/v3"


# --------------------------------------------------------------------------- #
# OAuth / pacing helpers (mirrors writeoff_invoices.py)
# --------------------------------------------------------------------------- #
# Zoho rate-limits *token generation* (~a handful per few minutes per refresh
# token). Every process start otherwise burns one; cache the access token to disk
# and reuse it across invocations until it is close to expiry.
TOKEN_CACHE = os.path.join(RUNS_DIR, ".token_cache.json")


class TokenManager:
    def __init__(self):
        self.lock = threading.Lock()
        self.token = ""
        self.expires_at = 0.0
        self.last_refreshed = 0.0
        self._load_disk()

    def _load_disk(self):
        try:
            with open(TOKEN_CACHE, encoding="utf-8") as f:
                d = json.load(f)
            if d.get("access_token") and d.get("expires_at", 0) > time.time() + 120:
                self.token = d["access_token"]
                self.expires_at = d["expires_at"]
        except (OSError, ValueError):
            pass

    def _save_disk(self):
        try:
            with open(TOKEN_CACHE, "w", encoding="utf-8") as f:
                json.dump({"access_token": self.token, "expires_at": self.expires_at}, f)
        except OSError:
            pass

    def get_token(self, force_refresh: bool = False) -> str:
        with self.lock:
            now = time.time()
            should_refresh = (
                not self.token
                or now >= self.expires_at
                or (force_refresh and (now - self.last_refreshed > 10))
            )
            if should_refresh:
                if force_refresh:
                    self._load_disk()  # another process may have refreshed already
                    if self.token and now < self.expires_at:
                        return self.token
                last_err = ""
                for _ in range(3):
                    try:
                        resp = requests.post(TOKEN_URL, params={
                            "grant_type": "refresh_token",
                            "refresh_token": REFRESH_TOKEN,
                            "client_id": CLIENT_ID,
                            "client_secret": CLIENT_SECRET,
                        }, timeout=20)
                        data = resp.json()
                        if "access_token" in data:
                            self.token = data["access_token"]
                            self.expires_at = now + data.get("expires_in", 3600) - 180
                            self.last_refreshed = now
                            self._save_disk()
                            return self.token
                        last_err = f"{resp.status_code} {data.get('error_description', data)}"
                    except Exception as e:
                        last_err = str(e)
                    time.sleep(2)
                if not self.token:
                    raise RuntimeError(f"Token refresh failed after retries: {last_err}")
            return self.token


token_mgr = TokenManager()


class RateLimiter:
    """Shared across worker threads -- caps the aggregate call rate, not per-thread."""

    def __init__(self, calls_per_min: int):
        self.interval = 60.0 / calls_per_min
        self.lock = threading.Lock()
        self.next_slot = time.monotonic()

    def wait(self):
        with self.lock:
            now = time.monotonic()
            slot = max(now, self.next_slot)
            self.next_slot = slot + self.interval
        delay = slot - now
        if delay > 0:
            time.sleep(delay)


class QuotaTracker:
    def __init__(self, min_remaining: int = 100):
        self.lock = threading.Lock()
        self.min_remaining = min_remaining
        self.stop_requested = False
        self.last_remaining = None

    def record_headers(self, headers):
        rem = headers.get("X-Rate-Limit-Remaining")
        if rem is None:
            return
        try:
            rem_val = int(rem)
        except ValueError:
            return
        with self.lock:
            self.last_remaining = rem_val
            if rem_val <= self.min_remaining and not self.stop_requested:
                self.stop_requested = True
                print(f"\n[ALERT] Daily API buffer reached (remaining {rem_val} <= {self.min_remaining}). "
                      f"Gracefully stopping; remainder will be postponed.\n", flush=True)

    def should_stop(self) -> bool:
        with self.lock:
            return self.stop_requested


def _headers() -> dict:
    return {"Authorization": f"Zoho-oauthtoken {token_mgr.get_token()}",
            "Content-Type": "application/json"}


def _request(method: str, url: str, quota: QuotaTracker | None = None,
             params_extra: dict | None = None, **kw):
    """One HTTP call with a 401 refresh retry and 429 backoff. Returns (resp, data)."""
    params = {"organization_id": ORG_ID, **(params_extra or {})}
    for attempt in range(4):
        resp = requests.request(method, url, headers=_headers(),
                                params=params, timeout=40, **kw)
        if quota:
            quota.record_headers(resp.headers)
        if resp.status_code == 401 and attempt == 0:
            token_mgr.get_token(force_refresh=True)
            continue
        if resp.status_code == 429:
            time.sleep(5)
            continue
        data = resp.json() if resp.content else {}
        return resp, data
    return resp, (resp.json() if resp.content else {})


def _resolve(path: str) -> str:
    if not path or not isinstance(path, str):
        raise FileNotFoundError("Empty or invalid file path provided.")
    if os.path.isfile(path):
        return os.path.abspath(path)
    cand1 = os.path.join(RUNS_DIR, path)
    if os.path.isfile(cand1):
        return os.path.abspath(cand1)
    cand2 = os.path.join(RUNS_DIR, os.path.basename(path))
    if os.path.isfile(cand2):
        return os.path.abspath(cand2)
    script_dir = os.path.dirname(os.path.abspath(__file__))
    cand3 = os.path.join(script_dir, path)
    if os.path.isfile(cand3):
        return os.path.abspath(cand3)
    raise FileNotFoundError(f"Could not find file: {path}")


def confirm(prompt: str) -> bool:
    return input(f"{prompt} [y/N]: ").strip().lower() == "y"


def _ts() -> str:
    return time.strftime("%Y%m%dT%H%M%SZ", time.gmtime())


def _now() -> str:
    return time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())


class CsvLogger:
    """Append-only CSV written row-by-row (flushed every write) so a crash still
    leaves a complete record of every API call made so far."""

    def __init__(self, path: str, fieldnames: list[str]):
        self.path = path
        self.fieldnames = fieldnames
        self.lock = threading.Lock()
        self._fh = open(path, "w", newline="", encoding="utf-8-sig")
        self._w = csv.DictWriter(self._fh, fieldnames=fieldnames, delimiter="|", extrasaction="ignore")
        self._w.writeheader()
        self._fh.flush()

    def write(self, row: dict):
        with self.lock:
            self._w.writerow(row)
            self._fh.flush()

    def close(self):
        with self.lock:
            self._fh.close()


def ledger_name_for(current_salesman: str) -> str:
    name = re.sub(r"^MR\s+", "", (current_salesman or "").strip(), flags=re.I).strip()
    return f"{name} OLD BALANCE"


def _cell_date(v) -> str:
    """xlsx date cell -> 'YYYY-MM-DD' string (openpyxl gives datetime or str)."""
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    return str(v)[:10]


def _cell_num(v):
    return round(float(v), 2) if isinstance(v, (int, float)) else None


# --------------------------------------------------------------------------- #
# Step A -- build candidates from the xlsx
# --------------------------------------------------------------------------- #
def build_candidates() -> str:
    from openpyxl import load_workbook

    wb = load_workbook(INPUT_XLSX, data_only=True)
    ws = wb[XLSX_SHEET]
    rows = list(ws.iter_rows(values_only=True))[1:]

    seen: dict[str, dict] = {}
    dropped = 0
    for r in rows:
        inv = r[C_INV]
        inv_id = r[C_INVID]
        cust_id = r[C_CUSTID]
        bal = r[C_BAL]
        if inv is None and inv_id is None:
            continue
        if not inv_id or not cust_id or not isinstance(bal, (int, float)) or bal <= 0:
            dropped += 1
            print(f"  [DROP] {inv!r}  inv_id={inv_id!r} cust_id={cust_id!r} bal={bal!r}", flush=True)
            continue
        inv_id = str(inv_id)
        bal = round(float(bal), 2)
        if inv_id in seen:
            if abs(seen[inv_id]["balance"] - bal) > 0.005:
                print(f"  [DUP!] {inv} {inv_id} balance {seen[inv_id]['balance']} vs {bal} -- keeping first", flush=True)
            continue
        seen[inv_id] = {
            "invoice_number": str(inv),
            "customer_name": str(r[C_CUST]),
            "salesman": str(r[C_SALES] or "").strip(),
            "current_salesman": str(r[C_CURSALES]),
            "ledger_name": ledger_name_for(r[C_CURSALES]),
            "invoice_date": _cell_date(r[C_INVDATE]),
            "due_date": _cell_date(r[C_DUEDATE]),
            "age_days": int(r[C_AGE]) if isinstance(r[C_AGE], (int, float)) else None,
            "invoice_total": _cell_num(r[C_TOTAL]),
            "balance": bal,
            "customer_id": str(cust_id),
            "invoice_id": inv_id,
        }

    candidates = list(seen.values())
    total = round(sum(c["balance"] for c in candidates), 2)

    by_ledger = defaultdict(lambda: [0, 0.0])
    for c in candidates:
        by_ledger[c["ledger_name"]][0] += 1
        by_ledger[c["ledger_name"]][1] += c["balance"]

    ts = _ts()
    cand_path = os.path.join(RUNS_DIR, f"candidates_{ts}.json")
    with open(cand_path, "w", encoding="utf-8") as f:
        json.dump(candidates, f, indent=2, ensure_ascii=False)

    lines = [f"Source : {INPUT_XLSX}",
             f"Built  : {ts}",
             f"Invoices: {len(candidates)}   Dropped: {dropped}",
             f"Total balance: {total:.2f}   (expected {EXPECTED_TOTAL:.2f})",
             "",
             f"{'Ledger':<32}{'Count':>8}{'Amount':>16}"]
    for name in sorted(by_ledger):
        cnt, amt = by_ledger[name]
        lines.append(f"{name:<32}{cnt:>8}{amt:>16.2f}")
    summary = "\n".join(lines)
    with open(os.path.join(RUNS_DIR, f"summary_{ts}.txt"), "w", encoding="utf-8") as f:
        f.write(summary + "\n")

    print(summary, flush=True)
    if abs(total - EXPECTED_TOTAL) > TOTAL_TOLERANCE:
        print(f"\n[WARN] total {total:.2f} deviates from expected {EXPECTED_TOTAL:.2f} "
              f"by more than {TOTAL_TOLERANCE}", flush=True)
    print(f"\nCandidates written: {cand_path}", flush=True)
    return cand_path


# --------------------------------------------------------------------------- #
# Step B -- ensure the parent + per-salesman sub-accounts exist
# --------------------------------------------------------------------------- #
def fetch_all_accounts() -> list[dict]:
    out, page = [], 1
    while True:
        resp = requests.get(f"{API_BASE}/chartofaccounts", headers=_headers(),
                            params={"organization_id": ORG_ID, "page": page, "per_page": 200},
                            timeout=40)
        data = resp.json() if resp.content else {}
        if resp.status_code != 200:
            raise RuntimeError(f"chartofaccounts fetch failed: {data}")
        out.extend(data.get("chartofaccounts", []))
        if not data.get("page_context", {}).get("has_more_page"):
            break
        page += 1
    return out


def resolve_ar_account_id() -> str:
    """The org's Accounts Receivable control account (credit side of every journal)."""
    global AR_ACCOUNT_ID
    if AR_ACCOUNT_ID:
        return AR_ACCOUNT_ID
    ar = [a for a in fetch_all_accounts() if a.get("account_type") == "accounts_receivable"]
    if len(ar) != 1:
        raise RuntimeError(f"expected exactly one accounts_receivable account, found {len(ar)}: "
                           f"{[a.get('account_name') for a in ar]}")
    AR_ACCOUNT_ID = ar[0]["account_id"]
    print(f"  A/R account: {ar[0]['account_name']} ({AR_ACCOUNT_ID})", flush=True)
    return AR_ACCOUNT_ID


def create_account(name: str, parent_id: str | None, account_code: str | None = None) -> dict:
    body = {"account_name": name, "account_type": ACCOUNT_TYPE}
    if account_code:
        body["account_code"] = account_code
    if parent_id:
        body["parent_account_id"] = parent_id
    resp, data = _request("POST", f"{API_BASE}/chartofaccounts", data=json.dumps(body))
    if resp.status_code not in (200, 201):
        raise RuntimeError(f"create account '{name}' failed: HTTP {resp.status_code} {data}")
    return data["chart_of_account"]


def ensure_ledgers(candidates: list[dict], allow_reuse: bool, ts: str) -> dict:
    wanted = sorted({c["ledger_name"] for c in candidates})
    all_accs = fetch_all_accounts()
    existing = {a["account_name"]: a for a in all_accs}
    used_codes = {a.get("account_code") for a in all_accs if a.get("account_code")}
    acc_csv = CsvLogger(os.path.join(RUNS_DIR, f"accounts_log_{ts}.csv"),
                        ["timestamp", "account_name", "action", "account_id", "account_code", "parent_account_id"])

    code_counter = SUBACCOUNT_CODE_BASE

    def next_free_code() -> str:
        nonlocal code_counter
        while True:
            cand = f"1.8.{code_counter}"
            code_counter += 100
            if cand not in used_codes:
                used_codes.add(cand)
                return cand

    def resolve(name: str, parent_id: str | None, default_code: str | None) -> str:
        if name in existing:
            acc_id = existing[name]["account_id"]
            code = existing[name].get("account_code") or default_code
            if allow_reuse:
                print(f"  [reuse]  {name}  ({acc_id})  code={code}", flush=True)
                acc_csv.write({"timestamp": _now(), "account_name": name, "action": "reuse",
                               "account_id": acc_id, "account_code": code or "",
                               "parent_account_id": parent_id or ""})
                return acc_id
            acc_csv.close()
            raise SystemExit(
                f"\n[STOP] Account '{name}' already exists (id {acc_id}).\n"
                f"       Plan says create fresh. Re-run with --allow-reuse to use it, "
                f"or rename PARENT_ACCOUNT / the sheet's salesman values.")
        assigned_code = default_code if (default_code and default_code not in used_codes) else next_free_code()
        used_codes.add(assigned_code)
        acc = create_account(name, parent_id, assigned_code)
        print(f"  [create] {name}  ({acc['account_id']})  code={assigned_code}"
              + (f"  parent={parent_id}" if parent_id else ""), flush=True)
        acc_csv.write({"timestamp": _now(), "account_name": name, "action": "create",
                       "account_id": acc["account_id"], "account_code": assigned_code or "",
                       "parent_account_id": parent_id or ""})
        existing[name] = acc
        return acc["account_id"]

    try:
        parent_id = resolve(PARENT_ACCOUNT, None, PARENT_ACCOUNT_CODE)
        ledger_map = {"__parent__": parent_id, "__parent_name__": PARENT_ACCOUNT}
        for name in wanted:
            ledger_map[name] = resolve(name, parent_id, None)
    finally:
        acc_csv.close()

    path = os.path.join(RUNS_DIR, f"ledger_map_{ts}.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(ledger_map, f, indent=2, ensure_ascii=False)
    print(f"\nLedger map written: {path}", flush=True)
    print(f"Accounts CSV      : {os.path.join(RUNS_DIR, f'accounts_log_{ts}.csv')}", flush=True)
    return ledger_map


# --------------------------------------------------------------------------- #
# Step C -- one manual journal per invoice, credit applied to that invoice
# --------------------------------------------------------------------------- #
def get_invoice_live(invoice_id: str, quota: QuotaTracker | None):
    resp, data = _request("GET", f"{API_BASE}/invoices/{invoice_id}", quota=quota)
    if resp.status_code != 200:
        return None, f"invoice fetch HTTP {resp.status_code}: {data.get('message', '')}"
    inv = data.get("invoice", {})
    return {"status": inv.get("status"), "balance": round(float(inv.get("balance", 0) or 0), 2)}, None


def existing_journal_id(c: dict, quota: QuotaTracker | None) -> str | None:
    """Idempotency guard: has this invoice already been settled by a prior run?
    Journals carry reference_number = OLDBAL-JC-<invoice_number>."""
    ref = f"{REF_PREFIX}-{c['invoice_number']}"
    resp, data = _request("GET", f"{API_BASE}/journals", quota=quota,
                          params_extra={"reference_number": ref})
    if resp.status_code != 200:
        return None
    for j in data.get("journals", []):
        if j.get("reference_number") == ref:
            return j.get("journal_id")
    return None


NOTES_MAX = 500  # Zoho hard limit on the journal Notes field


def _clean(s) -> str:
    """Zoho rejects < > (and newlines) in Notes/description -> strip them."""
    return re.sub(r"[<>\r\n]+", " ", str(s or "")).strip()


def journal_notes(c: dict, amount: float) -> str:
    """Human-readable story of the transfer, for an accountant reading the journal
    later without the spreadsheet. Kept under NOTES_MAX; lines with missing data
    are dropped."""
    inv = c["invoice_number"]
    L = [f"Aged-debt accountability transfer (old_bills.xlsx, {JOURNAL_DATE})."]

    dated = f"Invoice {inv}"
    if c.get("invoice_date"):
        dated += f" dated {c['invoice_date']}"
    if c.get("invoice_total") is not None:
        dated += f", total AED {c['invoice_total']:.2f}"
    if c.get("age_days") is not None:
        dated += f", {c['age_days']}d overdue"
    L.append(dated + ".")

    sm = _clean(c.get("salesman"))
    cur = _clean(c.get("current_salesman"))
    who = f" Salesman {sm}, now {cur}." if sm and cur and sm != cur else (
        f" Current salesman {cur}." if cur else "")
    L.append(f"Customer: {_clean(c['customer_name'])}.{who}")

    L.append(f"AED {amount:.2f} open balance moved from A/R to {c['ledger_name']} "
             f"(Other Current Asset).")
    L.append("Not a write-off; original revenue and VAT unchanged.")

    notes = " ".join(L)   # single line; Zoho rejects newlines / < > in Notes
    if len(notes) > NOTES_MAX:                 # last-resort guard for very long names
        notes = notes[:NOTES_MAX - 3].rstrip() + "..."
    return notes


def journal_body(c: dict, ledger_account_id: str, amount: float) -> dict:
    inv = c["invoice_number"]
    age = c.get("age_days")
    dbg = (f"salesman {_clean(c.get('salesman')) or 'n/a'}, "
           f"invoice {c.get('invoice_date') or 'n/a'}"
           + (f", {age}d overdue" if age is not None else ""))
    cust = _clean(c["customer_name"])
    return {
        "journal_date": JOURNAL_DATE,
        "reference_number": f"{REF_PREFIX}-{inv}",
        "notes": journal_notes(c, amount),
        "journal_type": "both",
        "line_items": [
            {"account_id": resolve_ar_account_id(), "customer_id": c["customer_id"],
             "debit_or_credit": "credit", "amount": amount,
             "description": (f"{inv} {cust} - aged open balance {amount:.2f} "
                             f"cleared from A/R to {c['ledger_name']}")},
            {"account_id": ledger_account_id,
             "debit_or_credit": "debit", "amount": amount,
             "description": (f"{inv} aged receivable parked in {c['ledger_name']} - {dbg}")},
        ],
    }


def _ar_credit_line_id(journal: dict, customer_id: str) -> str | None:
    for li in journal.get("line_items", []):
        if (li.get("debit_or_credit") == "credit"
                and str(li.get("customer_id") or "") == str(customer_id)):
            return li["line_id"]
    return None


def record_journal_transfer(c: dict, ledger_account_id: str, amount: float,
                            quota: QuotaTracker | None):
    """Returns (ok: bool, result_dict, message). result_dict carries journal_id /
    journal_line_id / entry_number - partially populated on a mid-way failure so
    --undo can still clean up."""
    out: dict = {}

    # 1. create (org auto-publishes; publish explicitly only if it didn't)
    resp, data = _request("POST", f"{API_BASE}/journals",
                          data=json.dumps(journal_body(c, ledger_account_id, amount)), quota=quota)
    if resp.status_code not in (200, 201):
        return False, out, f"journal create HTTP {resp.status_code}: {data.get('message', data)}"
    jrn = data.get("journal", {})
    jid = jrn.get("journal_id")
    out["journal_id"] = jid
    out["entry_number"] = jrn.get("entry_number")

    if jrn.get("status") != "published":
        resp, data = _request("POST", f"{API_BASE}/journals/{jid}/status/publish", quota=quota)
        if resp.status_code not in (200, 201) and data.get("code") != 1049:
            return False, out, f"publish HTTP {resp.status_code}: {data.get('message', data)}"

    # 2. find the receivables credit line
    resp, data = _request("GET", f"{API_BASE}/journals/{jid}/credits", quota=quota)
    line_id = None
    for cr in data.get("available_journal_credits", []):
        if (cr.get("is_receivable_credit")
                and str(cr.get("customer_id") or "") == str(c["customer_id"])
                and round(float(cr.get("available_credits", 0) or 0), 2) > 0):
            line_id = cr.get("journal_line_id")
            break
    if not line_id:  # fall back to the created journal's own line list
        line_id = _ar_credit_line_id(jrn, c["customer_id"])
    if not line_id:
        return False, out, "could not locate AR credit line on the published journal"
    out["journal_line_id"] = line_id

    # 3. apply the credit to the invoice
    resp, data = _request(
        "POST", f"{API_BASE}/journals/{jid}/credits/{line_id}/invoices",
        data=json.dumps({"invoices": [{"invoice_id": c["invoice_id"], "amount_applied": amount}]}),
        quota=quota)
    if resp.status_code not in (200, 201) or data.get("code") not in (0, None):
        return False, out, f"apply HTTP {resp.status_code}: {data.get('message', data)}"

    # 4. verify the invoice actually closed (proof, not just a 200 on the apply)
    live, err = get_invoice_live(c["invoice_id"], quota)
    if err:
        out["invoice_status_after"] = "?"
        out["residual_balance"] = None
    else:
        out["invoice_status_after"] = live["status"]
        out["residual_balance"] = live["balance"]
    return True, out, data.get("message", "OK")


def process_one(c: dict, ledger_map: dict, verify_balance: bool, quota: QuotaTracker) -> dict:
    base = dict(c)
    if quota.should_stop():
        return {**base, "success": False, "skipped": True, "message": "postponed: daily API buffer"}

    account_id = ledger_map.get(c["ledger_name"])
    if not account_id:
        return {**base, "success": False, "message": f"no ledger id for {c['ledger_name']}"}

    apply_amount = c["balance"]
    warning = None
    if verify_balance:
        live, err = get_invoice_live(c["invoice_id"], quota)
        if err:
            return {**base, "success": False, "message": err}
        if live["status"] in ("paid", "void", "draft") or live["balance"] <= 0:
            return {**base, "success": False, "skipped": True,
                    "message": f"skip: invoice status={live['status']} balance={live['balance']}"}
        apply_amount = round(min(c["balance"], live["balance"]), 2)
        if abs(apply_amount - c["balance"]) > 0.005:
            warning = f"sheet {c['balance']} -> applied {apply_amount} (live {live['balance']})"

    dup = existing_journal_id(c, quota)
    if dup:
        return {**base, "success": False, "skipped": True, "journal_id": dup,
                "amount_applied": apply_amount,
                "message": f"skip: journal {dup} already exists for this invoice"}

    ok, res, message = record_journal_transfer(c, account_id, apply_amount, quota)
    residual = res.get("residual_balance")
    entry = {**base, "success": ok,
             "journal_id": res.get("journal_id"), "journal_line_id": res.get("journal_line_id"),
             "entry_number": res.get("entry_number"),
             "amount_applied": apply_amount, "message": message,
             "invoice_status_after": res.get("invoice_status_after"),
             "residual_balance": residual}
    warns = [warning] if warning else []
    if ok and residual is not None and residual > 0.01:
        warns.append(f"invoice NOT fully closed - residual balance {residual:.2f} "
                     f"(status {res.get('invoice_status_after')})")
    if ok and residual is None:
        warns.append("post-apply verify GET failed - closure unconfirmed")
    if warns:
        entry["warning"] = "; ".join(warns)
    return entry


def run_settle(input_path: str, limit: int | None, calls_per_min: int, concurrency: int,
               min_remaining: int, verify_balance: bool, allow_reuse: bool,
               ledger_map_path: str | None, dry_run: bool, skip_confirm: bool):
    with open(_resolve(input_path), encoding="utf-8") as f:
        candidates = json.load(f)
    if limit is not None:
        candidates = candidates[:limit]
    if not candidates:
        print("Nothing to do.", flush=True)
        return

    total = round(sum(c["balance"] for c in candidates), 2)
    by_ledger = defaultdict(lambda: [0, 0.0])
    for c in candidates:
        by_ledger[c["ledger_name"]][0] += 1
        by_ledger[c["ledger_name"]][1] += c["balance"]

    print(f"Loaded {len(candidates)} invoice(s) from {input_path}", flush=True)
    print(f"Total sheet balance: {total:.2f}", flush=True)
    for name in sorted(by_ledger):
        cnt, amt = by_ledger[name]
        print(f"  {name:<32}{cnt:>7}{amt:>15.2f}", flush=True)

    token_mgr.get_token()

    if dry_run:
        print("\n--- DRY RUN: no writes ---", flush=True)
        try:
            ar_id = resolve_ar_account_id()
        except Exception as e:
            ar_id = f"<accounts_receivable: {e}>"
        for c in candidates[:3]:
            sample = journal_body(c, f"<{c['ledger_name']}>", c["balance"])
            sample["line_items"][0]["account_id"] = ar_id
            print(f"\n  POST {API_BASE}/journals", flush=True)
            print("  notes:\n    " + sample["notes"].replace("\n", "\n    "), flush=True)
            print(json.dumps(sample, indent=2), flush=True)
            print(f"  then GET  {API_BASE}/journals/{{id}}/credits", flush=True)
            print(f"  then POST {API_BASE}/journals/{{id}}/credits/{{line_id}}/invoices "
                  + json.dumps({"invoices": [{"invoice_id": c["invoice_id"],
                                              "amount_applied": c["balance"]}]}), flush=True)
        print("\n  account plan:", flush=True)
        try:
            existing = {a["account_name"]: a["account_id"] for a in fetch_all_accounts()}
            for name in [PARENT_ACCOUNT] + sorted(by_ledger):
                state = f"EXISTS ({existing[name]})" if name in existing else "will CREATE"
                print(f"    {name:<34} {state}", flush=True)
        except Exception as e:
            print(f"    [could not read chart of accounts: {e}]", flush=True)
        return

    ts = _ts()
    if ledger_map_path:
        with open(_resolve(ledger_map_path), encoding="utf-8") as f:
            ledger_map = json.load(f)
        print(f"\nUsing existing ledger map: {ledger_map_path}", flush=True)
    else:
        print("\nEnsuring ledgers...", flush=True)
        ledger_map = ensure_ledgers(candidates, allow_reuse, ts)

    resolve_ar_account_id()  # fail fast before spinning up workers

    if not skip_confirm and not confirm(
            f"\nPost {len(candidates)} manual journal(s) totalling ~{total:.2f}, "
            f"crediting A/R and applying each to its invoice?"):
        print("Aborted. No journals posted.", flush=True)
        return

    limiter = RateLimiter(calls_per_min)
    quota = QuotaTracker(min_remaining=min_remaining)
    io_lock = threading.Lock()
    log: list[dict] = []
    ok = 0
    done = 0
    processed_ids: set[str] = set()

    jsonl_path = os.path.join(RUNS_DIR, f"settle_log_{ts}.jsonl")
    jsonl = open(jsonl_path, "a", encoding="utf-8")
    csv_path = os.path.join(RUNS_DIR, f"settle_log_{ts}.csv")
    csv_log = CsvLogger(csv_path, [
        "timestamp", "seq", "status", "invoice_number", "customer_name",
        "salesman", "current_salesman", "ledger_name", "invoice_date", "age_days",
        "customer_id", "invoice_id", "sheet_balance", "amount_applied",
        "journal_id", "journal_line_id", "entry_number",
        "invoice_status_after", "residual_balance", "message", "warning",
    ])
    print(f"Live log : {jsonl_path}", flush=True)
    print(f"CSV log  : {csv_path}\n", flush=True)

    def task(c: dict) -> dict:
        if quota.should_stop():
            return {**c, "success": False, "skipped": True, "message": "postponed: daily API buffer"}
        limiter.wait()
        try:
            return process_one(c, ledger_map, verify_balance, quota)
        except Exception as e:
            return {**c, "success": False, "message": f"exception: {e}"}

    try:
        with ThreadPoolExecutor(max_workers=concurrency) as pool:
            futures = {pool.submit(task, c): c for c in candidates}
            for fut in as_completed(futures):
                entry = fut.result()
                done += 1
                processed_ids.add(entry["invoice_id"])
                if entry["success"]:
                    status = "WARN" if entry.get("warning") else "OK  "
                elif entry.get("skipped"):
                    status = "SKIP"
                else:
                    status = "FAIL"
                with io_lock:
                    tail = f"{entry['message']}"
                    if entry.get("warning"):
                        tail += f"  [!] {entry['warning']}"
                    print(f"  [{done}/{len(candidates)}] [{status.strip():<4}] "
                          f"{entry['customer_name'][:34]:<34} {entry['invoice_number']:<16} "
                          f"{tail}", flush=True)
                    jsonl.write(json.dumps(entry, ensure_ascii=False) + "\n")
                    jsonl.flush()
                csv_log.write({
                    "timestamp": _now(), "seq": done, "status": status.strip(),
                    "invoice_number": entry["invoice_number"], "customer_name": entry["customer_name"],
                    "salesman": entry.get("salesman", ""),
                    "current_salesman": entry["current_salesman"], "ledger_name": entry["ledger_name"],
                    "invoice_date": entry.get("invoice_date", ""), "age_days": entry.get("age_days", ""),
                    "customer_id": entry["customer_id"], "invoice_id": entry["invoice_id"],
                    "sheet_balance": entry["balance"], "amount_applied": entry.get("amount_applied", ""),
                    "journal_id": entry.get("journal_id", ""),
                    "journal_line_id": entry.get("journal_line_id", ""),
                    "entry_number": entry.get("entry_number", ""),
                    "invoice_status_after": entry.get("invoice_status_after", ""),
                    "residual_balance": entry.get("residual_balance", ""),
                    "message": entry["message"], "warning": entry.get("warning", ""),
                })
                log.append(entry)
                if entry["success"]:
                    ok += 1
    finally:
        jsonl.close()
        csv_log.close()

    log_path = os.path.join(RUNS_DIR, f"settle_log_{ts}.json")
    with open(log_path, "w", encoding="utf-8") as f:
        json.dump(log, f, indent=2, ensure_ascii=False)

    applied = round(sum(e.get("amount_applied", 0) for e in log if e["success"]), 2)
    per_ledger = defaultdict(float)
    for e in log:
        if e["success"]:
            per_ledger[e["ledger_name"]] += e.get("amount_applied", 0)

    closed = [e for e in log if e["success"] and e.get("residual_balance") == 0]
    residual = [e for e in log if e["success"] and (e.get("residual_balance") or 0) > 0.01]
    unconfirmed = [e for e in log if e["success"] and e.get("residual_balance") is None]

    print(f"\n{ok}/{done} journals posted & applied. Applied total: {applied:.2f}", flush=True)
    for name in sorted(per_ledger):
        print(f"  {name:<32}{per_ledger[name]:>15.2f}", flush=True)
    print(f"\n{len(closed)} invoice(s) confirmed closed (balance 0); "
          f"{len(residual)} with residual balance; {len(unconfirmed)} unverified.", flush=True)
    for e in residual:
        print(f"  [RESIDUAL] {e['invoice_number']:<16} {e.get('residual_balance')} "
              f"(journal {e.get('entry_number')})", flush=True)
    for e in unconfirmed:
        print(f"  [UNVERIFIED] {e['invoice_number']:<16} (journal {e.get('entry_number')})", flush=True)
    print(f"JSON log: {log_path}", flush=True)
    print(f"CSV log : {csv_path}", flush=True)

    _CAND_KEYS = ("invoice_number", "customer_name", "salesman", "current_salesman",
                  "ledger_name", "invoice_date", "due_date", "age_days", "invoice_total",
                  "balance", "customer_id", "invoice_id")
    postponed = [c for c in candidates if c["invoice_id"] not in processed_ids] + \
                [e for e in log if e.get("skipped") and "postponed" in e["message"]]
    if postponed:
        pp = os.path.join(RUNS_DIR, f"candidates_postponed_{ts}.json")
        with open(pp, "w", encoding="utf-8") as f:
            json.dump([{k: c.get(k) for k in _CAND_KEYS} for c in postponed],
                      f, indent=2, ensure_ascii=False)
        print(f"\n[POSTPONED] {len(postponed)} invoice(s): {pp}", flush=True)


# --------------------------------------------------------------------------- #
# Step D -- undo
# --------------------------------------------------------------------------- #
def run_undo(log_path: str, undo_accounts: bool, ledger_map_path: str | None):
    with open(_resolve(log_path), encoding="utf-8") as f:
        entries = json.load(f)
    # anything that got a journal_id, even a partial failure, so a half-done row cleans up too
    succeeded = [e for e in entries if e.get("journal_id")]
    if not succeeded:
        print("No journals in this log to undo.", flush=True)
        return

    print(f"{len(succeeded)} journal(s) to unapply + delete from {log_path}", flush=True)
    if not confirm("Delete all of them (re-opens the invoices)?"):
        print("Aborted.", flush=True)
        return

    token_mgr.get_token()
    ts = _ts()
    csv_path = os.path.join(RUNS_DIR, f"undo_log_{ts}.csv")
    csv_log = CsvLogger(csv_path, [
        "timestamp", "seq", "kind", "target", "invoice_number", "journal_id", "result", "message",
    ])
    out: list[dict] = []
    ok = 0
    try:
        for i, e in enumerate(succeeded, 1):
            jid = e["journal_id"]
            # unapply every credit this journal placed on an invoice. The resource
            # id is invoices_credited[].journal_invoice_id (NOT journal_line_id,
            # despite the doc's parameter name) -> 404 otherwise.
            r, d = _request("GET", f"{API_BASE}/journals/{jid}")
            for jinv in d.get("journal", {}).get("invoices_credited", []) or []:
                jinv_id = jinv.get("journal_invoice_id")
                if not jinv_id:
                    continue
                r2, d2 = _request(
                    "DELETE", f"{API_BASE}/journals/{jid}/credits/{jinv_id}/receivables")
                csv_log.write({"timestamp": _now(), "seq": i, "kind": "unapply_credit",
                               "target": f"{jid}/{jinv_id}", "invoice_number": e["invoice_number"],
                               "journal_id": jid, "result": "OK" if r2.status_code == 200 else "SKIP",
                               "message": d2.get("message", str(r2.status_code))})
            resp, data = _request("DELETE", f"{API_BASE}/journals/{jid}")
            good = resp.status_code == 200
            ok += good
            msg = data.get("message", str(resp.status_code))
            print(f"  [{'OK' if good else 'FAIL'}] {e['invoice_number']:<16} {msg}", flush=True)
            out.append({**e, "undo_success": good, "undo_message": msg})
            csv_log.write({"timestamp": _now(), "seq": i, "kind": "delete_journal",
                           "target": jid, "invoice_number": e["invoice_number"],
                           "journal_id": jid, "result": "OK" if good else "FAIL",
                           "message": msg})
            time.sleep(0.75)
        print(f"\n{ok}/{len(succeeded)} journals deleted.", flush=True)

        if undo_accounts and ledger_map_path:
            with open(_resolve(ledger_map_path), encoding="utf-8") as f:
                ledger_map = json.load(f)
            ids = [(k, v) for k, v in ledger_map.items() if not k.startswith("__")]
            ids.append((ledger_map.get("__parent_name__", PARENT_ACCOUNT), ledger_map["__parent__"]))
            for j, (name, acc_id) in enumerate(ids, 1):
                resp, data = _request("POST", f"{API_BASE}/chartofaccounts/{acc_id}/inactive")
                good = resp.status_code == 200
                msg = data.get("message", str(resp.status_code))
                print(f"  [{'OK' if good else 'SKIP'}] deactivate {name}: {msg}", flush=True)
                csv_log.write({"timestamp": _now(), "seq": j, "kind": "deactivate_account",
                               "target": acc_id, "invoice_number": "", "journal_id": "",
                               "result": "OK" if good else "SKIP", "message": f"{name}: {msg}"})
    finally:
        csv_log.close()

    undo_path = os.path.join(RUNS_DIR, f"undo_log_{ts}.json")
    with open(undo_path, "w", encoding="utf-8") as f:
        json.dump(out, f, indent=2, ensure_ascii=False)
    print(f"Undo JSON log: {undo_path}", flush=True)
    print(f"Undo CSV log : {csv_path}", flush=True)


# --------------------------------------------------------------------------- #
def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    g = p.add_mutually_exclusive_group(required=True)
    g.add_argument("--build", action="store_true", help="Build candidates JSON from the xlsx (no writes)")
    g.add_argument("--input", help="candidates_<ts>.json to process")
    g.add_argument("--undo", help="settle_log_<ts>.json to reverse (unapply credit + delete each journal)")

    p.add_argument("--dry-run", action="store_true", help="With --input: print plan + sample journal payloads, no writes")
    p.add_argument("--limit", type=int, default=None, help="Process only the first N candidates")
    p.add_argument("--calls-per-min", type=int, default=80)
    p.add_argument("--concurrency", type=int, default=1)
    p.add_argument("--min-remaining", type=int, default=100, help="Stop when Zoho daily quota buffer hits this")
    p.add_argument("--no-verify-balance", action="store_true", help="Skip the per-invoice live balance GET")
    p.add_argument("--allow-reuse", action="store_true", help="Reuse a same-named existing account instead of stopping")
    p.add_argument("--ledger-map", default=None, help="Use an existing ledger_map_<ts>.json instead of creating accounts")
    p.add_argument("--undo-accounts", action="store_true", help="With --undo: also deactivate the created accounts")
    p.add_argument("--yes", action="store_true", help="Skip the confirmation prompt")
    args = p.parse_args()

    if args.build:
        build_candidates()
    elif args.input:
        run_settle(args.input, limit=args.limit, calls_per_min=args.calls_per_min,
                   concurrency=args.concurrency, min_remaining=args.min_remaining,
                   verify_balance=not args.no_verify_balance, allow_reuse=args.allow_reuse,
                   ledger_map_path=args.ledger_map, dry_run=args.dry_run, skip_confirm=args.yes)
    else:
        run_undo(args.undo, undo_accounts=args.undo_accounts, ledger_map_path=args.ledger_map)


if __name__ == "__main__":
    main()
