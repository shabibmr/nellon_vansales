"""
Find Write-off Candidates
--------------------------
Fetches all open invoices from Zoho Books, fully paginated, and produces an
Excel workbook with two sheets:
  - "All Customers"       : every customer with an outstanding balance, one row each (aggregated).
  - "Writeoff Candidates" : individual invoices matching a two-tier age/balance rule (see below) —
                            the ones actually worth writing off.

Two-tier candidate rule (by invoice date age):
  - age > OLD_AGE_DAYS (default 30)                      -> balance < OLD_MAX_BALANCE (default 30)
  - MID_MIN_AGE_DAYS (default 15) < age <= OLD_AGE_DAYS   -> balance < MID_MAX_BALANCE (default 15)
  - age <= MID_MIN_AGE_DAYS                                -> never a candidate, regardless of balance

Also writes the candidates as JSON (writeoff_invoices.py consumes this — Excel is for human review).
Read-only — does not change anything in Zoho.

Usage:
    pip install requests python-dotenv openpyxl
    python find_writeoff_candidates.py
    python find_writeoff_candidates.py --old-age-days 30 --old-max-balance 30 --mid-min-age-days 15 --mid-max-balance 15
"""

import argparse
import json
import os
import time
from datetime import date
import requests
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

load_dotenv(os.path.join(os.path.dirname(__file__), "..", ".env_zoho"))


def _require_env(key: str) -> str:
    val = os.getenv(key)
    if not val:
        raise RuntimeError(f"Missing required env var: {key}  (set it in .env_zoho)")
    return val


CLIENT_ID = _require_env("ZOHO_CLIENT_ID")
CLIENT_SECRET = _require_env("ZOHO_CLIENT_SECRET")
REFRESH_TOKEN = _require_env("ZOHO_REFRESH_TOKEN")
ORG_ID = _require_env("ZOHO_ORG_ID")
TOKEN_URL = "https://accounts.zoho.com/oauth/v2/token"
API_BASE = "https://www.zohoapis.com/books/v3"

RUNS_DIR = os.path.join(os.path.dirname(__file__), "writeoff_runs")

# Statuses that can carry a real outstanding balance worth writing off.
# draft/void invoices are excluded — nothing genuinely owed there.
RELEVANT_STATUSES = {"sent", "overdue", "unpaid", "partially_paid", "viewed"}


def get_access_token() -> str:
    resp = requests.post(TOKEN_URL, params={
        "grant_type": "refresh_token",
        "refresh_token": REFRESH_TOKEN,
        "client_id": CLIENT_ID,
        "client_secret": CLIENT_SECRET,
    })
    data = resp.json()
    if "access_token" not in data:
        raise RuntimeError(f"Token request failed: {data}")
    return data["access_token"]


def get_currency_symbol(token: str) -> str:
    headers = {"Authorization": f"Zoho-oauthtoken {token}"}
    resp = requests.get(f"{API_BASE}/organizations/{ORG_ID}", headers=headers, timeout=30)
    if resp.status_code != 200:
        return ""
    return resp.json().get("organization", {}).get("currency_code", "")


def fetch_all_invoices(token: str) -> list[dict]:
    """Fetch every non-paid/void/draft invoice across all pages.
    status=unpaid is a server-side aggregate filter (confirmed against the live org) that
    already excludes paid/void/draft — far fewer pages than pulling the entire all-time history."""
    headers = {
        "Authorization": f"Zoho-oauthtoken {token}",
        "Content-Type": "application/json",
    }
    all_invoices = []
    page = 1
    while True:
        params = {"organization_id": ORG_ID, "per_page": 200, "page": page, "status": "unpaid"}
        resp = requests.get(f"{API_BASE}/invoices", headers=headers, params=params, timeout=30)
        data = resp.json()
        if resp.status_code != 200:
            raise RuntimeError(f"Fetch failed on page {page}: HTTP {resp.status_code} {data.get('message', '')}")

        invoices = data.get("invoices", [])
        all_invoices.extend(invoices)

        page_context = data.get("page_context", {})
        print(f"  [OK] page {page} -> {len(invoices)} invoices "
              f"(total so far: {len(all_invoices)})")

        if not page_context.get("has_more_page"):
            break
        page += 1
        time.sleep(0.3)  # be gentle on rate limits

    return all_invoices


def _invoice_age_days(invoice: dict, today: date) -> int | None:
    raw = invoice.get("date")
    if not raw:
        return None
    try:
        invoice_date = date.fromisoformat(raw)
    except ValueError:
        return None
    return (today - invoice_date).days


def _candidate_tier(balance: float, age: int | None, args) -> str | None:
    """Return which tier rule this invoice matches, or None if it's not a candidate."""
    if age is None:
        return None
    if age > args.old_age_days and balance < args.old_max_balance:
        return "old"
    if args.mid_min_age_days < age <= args.old_age_days and balance < args.mid_max_balance:
        return "mid"
    return None


# User-requested exclusion: "CASH CUSTOMER" invoices dated in 2026 with balance > 10 AED
# are excluded from candidates even if they'd otherwise match a tier — these are being
# reviewed separately, not written off in this batch.
CASH_CUSTOMER_EXCLUDE_YEAR = 2026
CASH_CUSTOMER_EXCLUDE_MIN_BALANCE = 10


def _is_excluded(invoice: dict, balance: float) -> bool:
    if (invoice.get("customer_name") or "").strip().upper() != "CASH CUSTOMER":
        return False
    raw = invoice.get("date")
    if not raw:
        return False
    try:
        invoice_year = date.fromisoformat(raw).year
    except ValueError:
        return False
    return invoice_year == CASH_CUSTOMER_EXCLUDE_YEAR and balance > CASH_CUSTOMER_EXCLUDE_MIN_BALANCE


def _autosize(ws):
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(length + 2, 45)


def build_workbook(open_invoices: list[dict], candidates: list[dict], currency: str, today: date) -> Workbook:
    wb = Workbook()
    bold = Font(bold=True)

    # --- Sheet 1: All Customers (aggregated) ---
    ws1 = wb.active
    ws1.title = "All Customers"
    headers1 = ["Customer", f"Total Balance ({currency})" if currency else "Total Balance",
                "Open Invoices", "Oldest Invoice Date", "Oldest Invoice Age (days)"]
    ws1.append(headers1)
    for cell in ws1[1]:
        cell.font = bold

    by_customer: dict[str, dict] = {}
    for inv in open_invoices:
        name = inv.get("customer_name", "Unknown")
        balance = float(inv.get("balance", 0) or 0)
        age = _invoice_age_days(inv, today)
        entry = by_customer.setdefault(name, {"total": 0.0, "count": 0, "oldest_date": None, "oldest_age": None})
        entry["total"] += balance
        entry["count"] += 1
        inv_date = inv.get("date", "")
        if entry["oldest_date"] is None or (inv_date and inv_date < entry["oldest_date"]):
            entry["oldest_date"] = inv_date
            entry["oldest_age"] = age

    for name, entry in sorted(by_customer.items(), key=lambda kv: -kv[1]["total"]):
        ws1.append([name, round(entry["total"], 2), entry["count"], entry["oldest_date"], entry["oldest_age"]])
    _autosize(ws1)

    # --- Sheet 2: Writeoff Candidates (per-invoice) ---
    ws2 = wb.create_sheet("Writeoff Candidates")
    headers2 = ["Customer", "Invoice #", "Status", "Invoice Date", "Age (days)",
                f"Balance ({currency})" if currency else "Balance", "Tier"]
    ws2.append(headers2)
    for cell in ws2[1]:
        cell.font = bold

    for inv in candidates:
        ws2.append([
            inv.get("customer_name", ""), inv.get("invoice_number", ""), inv.get("status", ""),
            inv.get("date", ""), inv.get("age_days"), round(float(inv["balance"]), 2), inv.get("tier", ""),
        ])
    _autosize(ws2)

    return wb


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--old-age-days", type=int, default=30,
                         help="Age threshold (days) above which the looser 'old' balance cap applies. Default 30.")
    parser.add_argument("--old-max-balance", type=float, default=30,
                         help="Balance cap (exclusive) for invoices older than --old-age-days. Default 30.")
    parser.add_argument("--mid-min-age-days", type=int, default=15,
                         help="Lower age bound (days, exclusive) for the stricter 'mid' tier. Default 15.")
    parser.add_argument("--mid-max-balance", type=float, default=15,
                         help="Balance cap (exclusive) for invoices between --mid-min-age-days and --old-age-days. Default 15.")
    args = parser.parse_args()

    os.makedirs(RUNS_DIR, exist_ok=True)
    today = date.today()

    print("Fetching OAuth access token ...")
    token = get_access_token()
    currency = get_currency_symbol(token)

    print("Fetching all invoices (paginated) ...")
    invoices = fetch_all_invoices(token)
    print(f"Total invoices fetched: {len(invoices)}\n")

    open_invoices = [inv for inv in invoices
                      if inv.get("status") in RELEVANT_STATUSES and float(inv.get("balance", 0) or 0) > 0]

    candidates_raw = []
    excluded_count = 0
    for inv in open_invoices:
        balance = float(inv.get("balance", 0) or 0)
        age = _invoice_age_days(inv, today)
        tier = _candidate_tier(balance, age, args)
        if not tier:
            continue
        if _is_excluded(inv, balance):
            excluded_count += 1
            continue
        candidates_raw.append({**inv, "age_days": age, "tier": tier})
    candidates_raw.sort(key=lambda inv: float(inv["balance"]))
    if excluded_count:
        print(f"Excluded {excluded_count} CASH CUSTOMER invoice(s) dated {CASH_CUSTOMER_EXCLUDE_YEAR} "
              f"with balance > {CASH_CUSTOMER_EXCLUDE_MIN_BALANCE} {currency}.\n")

    print(f"{len(open_invoices)} open invoice(s) across all customers.")
    print(f"Rule: age > {args.old_age_days}d -> balance < {args.old_max_balance} {currency}  |  "
          f"{args.mid_min_age_days}d < age <= {args.old_age_days}d -> balance < {args.mid_max_balance} {currency}")
    print(f"{len(candidates_raw)} candidate(s) matched.\n")

    if candidates_raw:
        print(f"{'Customer':<35} {'Invoice #':<18} {'Age(d)':>7} {'Balance':>10}  Tier")
        print("-" * 82)
        total = 0.0
        for inv in candidates_raw:
            balance = float(inv["balance"])
            total += balance
            print(f"{inv.get('customer_name', ''):<35} {inv.get('invoice_number', ''):<18} "
                  f"{inv['age_days']:>7} {balance:>10.2f}  {inv['tier']}")
        print("-" * 82)
        print(f"{len(candidates_raw)} invoice(s), total balance: {total:.2f}\n")

    timestamp = time.strftime("%Y%m%dT%H%M%SZ", time.gmtime())

    # Excel workbook — for human review
    wb = build_workbook(open_invoices, candidates_raw, currency, today)
    xlsx_path = os.path.join(RUNS_DIR, f"writeoff_review_{timestamp}.xlsx")
    wb.save(xlsx_path)
    print(f"Excel review file saved to: {xlsx_path}")

    if not candidates_raw:
        print("No candidates matched the filters — nothing to write off. See the Excel file for the full picture.")
        return

    # JSON — exact input for writeoff_invoices.py, must match the "Writeoff Candidates" sheet
    out = [
        {
            "invoice_id": inv["invoice_id"],
            "invoice_number": inv.get("invoice_number", ""),
            "customer_name": inv.get("customer_name", ""),
            "status": inv.get("status", ""),
            "date": inv.get("date", ""),
            "age_days": inv["age_days"],
            "balance": float(inv["balance"]),
            "tier": inv["tier"],
        }
        for inv in candidates_raw
    ]
    json_path = os.path.join(RUNS_DIR, f"candidates_{timestamp}.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(out, f, indent=2, ensure_ascii=False)

    print(f"Candidate list (for script 2) saved to: {json_path}")
    print("Nothing has been changed in Zoho. Review the Excel file, then run:")
    print(f"  python writeoff_invoices.py --input {os.path.basename(json_path)}")


if __name__ == "__main__":
    main()
