"""CSV / Excel helpers for MIS report rows."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any, BinaryIO, Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def write_csv(rows: Sequence[dict], columns: Sequence[str], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(columns), extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({c: row.get(c, "") for c in columns})


def write_excel_sheets(
    sheets: dict[str, tuple[Sequence[str], Iterable[dict]]],
    path: Path,
) -> None:
    """Write multiple sheets: name -> (columns, rows)."""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    first = True
    for sheet_name, (columns, rows) in sheets.items():
        title = sheet_name[:31]
        if first:
            ws = wb.active
            ws.title = title
            first = False
        else:
            ws = wb.create_sheet(title)
        ws.append(list(columns))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([row.get(c, "") for c in columns])
        _format_numbers(ws, columns)
        for i, _ in enumerate(columns, start=1):
            ws.column_dimensions[get_column_letter(i)].width = 18
    wb.save(path)


_NUM = "#,##0.00_);[Red](#,##0.00)"  # accounting: negatives red, in brackets
_PCT = "0.0%_);[Red](0.0%)"
_FILL_TOTAL = PatternFill("solid", fgColor="EEF2F7")
_FILL_CHECK_OK = PatternFill("solid", fgColor="E3F4E8")
_FILL_CHECK_BAD = PatternFill("solid", fgColor="FBE3E1")


_MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


MARGIN_TEXT_MAX = 50  # Margin sheets: text columns are 50 chars wide; longer values are cut with "…"


DATE_WIDTH = 12  # "2026-09-30" plus padding


def _is_margin(name: str) -> bool:
    return name.startswith("Margin")


def _ellide(v: Any, limit: int = MARGIN_TEXT_MAX) -> Any:
    return v[: limit - 1] + "…" if isinstance(v, str) and len(v) > limit else v


def _format_numbers(ws, columns: Sequence[str]) -> None:
    """Accounting format on numeric data cells (negatives red, in brackets); *_pct columns as percentages."""
    for ci, name in enumerate(columns, start=1):
        fmt = _PCT if str(name).endswith("pct") else _NUM
        for (cell,) in ws.iter_rows(min_row=2, min_col=ci, max_col=ci):
            if isinstance(cell.value, float):
                cell.number_format = fmt


def month_label(col: str) -> str:
    """'2026-04' -> 'Apr-26' (template style); anything else unchanged."""
    if len(col) == 7 and col[4] == "-" and col[:4].isdigit() and col[5:].isdigit():
        return f"{_MONTHS[int(col[5:]) - 1]}-{col[2:4]}"
    return col


def write_statement_workbook(
    statements: Sequence[Any],
    path: Path | BinaryIO,
    data_sheets: dict[str, tuple[Sequence[str], Iterable[dict]]] | None = None,
    sheet_order: Sequence[str] | None = None,
) -> None:
    """Formatted MIS workbook: one sheet per `reports.financials.Statement`
    (or anything with name/title/subtitle/columns/lines), then optional plain
    data sheets (name -> (columns, rows)). `path` may be a file object."""
    if isinstance(path, Path):
        path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)

    for st in statements:
        ws = wb.create_sheet(st.name[:31])
        ws.append([st.title])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([st.subtitle])
        ws["A2"].font = Font(italic=True, color="666666")
        ws.append([])
        ws.append([getattr(st, "first_header", "Particulars"), *[month_label(c) for c in st.columns]])
        for cell in ws[4]:
            cell.font = Font(bold=True)
            cell.border = Border(bottom=Side(style="thin"))
        for col in range(2, len(st.columns) + 2):
            ws.cell(row=4, column=col).alignment = Alignment(horizontal="right")

        for ln in st.lines:
            if ln.kind == "blank":
                ws.append([])
                continue
            values = [] if ln.kind == "title" else [ln.values.get(c, 0.0) for c in st.columns]
            ws.append([_ellide(ln.label) if _is_margin(st.name) else ln.label, *values])
            r = ws.max_row
            label = ws.cell(row=r, column=1)
            if ln.kind == "title":
                label.font = Font(bold=True, color="1F4E79")
                continue
            if ln.kind == "sub":
                label.alignment = Alignment(indent=2)
                label.font = Font(color="555555")
            elif ln.kind == "pct":
                label.font = Font(italic=True)
            elif ln.kind in ("total", "category"):
                label.font = Font(bold=ln.kind == "total")
            fill = None
            if ln.kind == "total":
                fill = _FILL_TOTAL
            elif ln.kind == "check":
                ok = all(abs(v) < 0.015 for v in ln.values.values())
                fill = _FILL_CHECK_OK if ok else _FILL_CHECK_BAD
            for col in range(1, len(st.columns) + 2):
                cell = ws.cell(row=r, column=col)
                if col > 1:
                    cell.number_format = _PCT if ln.kind == "pct" else _NUM
                    if ln.kind == "total":
                        cell.font = Font(bold=True)
                if fill:
                    cell.fill = fill

        ws.column_dimensions["A"].width = MARGIN_TEXT_MAX if _is_margin(st.name) else 44
        for col in range(2, len(st.columns) + 2):
            ws.column_dimensions[get_column_letter(col)].width = 16
        ws.freeze_panes = "B5"

    for sheet_name, (columns, rows) in (data_sheets or {}).items():
        ws = wb.create_sheet(sheet_name[:31])
        ws.append(list(columns))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        margin = _is_margin(sheet_name)
        for row in rows:
            ws.append([_ellide(row.get(c, "")) if margin else row.get(c, "") for c in columns])
        _format_numbers(ws, columns)
        for i, c in enumerate(columns, start=1):
            width = 18
            if margin:  # fit the longest text value (header included), capped at 50 chars
                longest = max([len(str(c))] + [len(str(cell.value)) for cell in ws[get_column_letter(i)][1:]
                                               if isinstance(cell.value, str)])
                width = max(width, min(longest + 2, MARGIN_TEXT_MAX))
            if margin and c == "date":
                width = DATE_WIDTH
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A2"

    if sheet_order:  # listed sheets first in that order; any others keep their relative order after
        rank = {n[:31]: i for i, n in enumerate(sheet_order)}
        wb._sheets.sort(key=lambda w: rank.get(w.title, len(rank)))

    wb.save(path)
