#!/usr/bin/env python3
"""
Export Trial Balance from Zoho Books via General Ledger + Chart of Accounts.

Manual /journals vouchers alone omit invoices, bills, expenses, and receipts.
This script builds a period trial balance from:

  GET /reports/generalledger
      ?filter_by=TransactionDate.CustomDate&from_date=&to_date=

  GET /chartofaccounts?showbalance=true

Credentials: ../../.env_zoho
  ZOHO_CLIENT_ID, ZOHO_CLIENT_SECRET, ZOHO_REFRESH_TOKEN, ZOHO_ORG_ID

Usage:
  .venv/bin/python -m scripts.export_trial_balance
  .venv/bin/python -m scripts.export_trial_balance --from-date 2026-04-01 --to-date 2026-10-11
  .venv/bin/python -m scripts.export_trial_balance --from-date 2026-04-01 --to-date 2026-10-11 --include-zero
"""

from __future__ import annotations

import argparse
import csv
import json
import os
from datetime import date, datetime
from pathlib import Path
from typing import Any

import requests
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[3]  # nellon_vansales/
load_dotenv(ROOT / ".env_zoho")

CLIENT_ID = os.environ["ZOHO_CLIENT_ID"]
CLIENT_SECRET = os.environ["ZOHO_CLIENT_SECRET"]
REFRESH_TOKEN = os.environ["ZOHO_REFRESH_TOKEN"]
ORG_ID = os.environ["ZOHO_ORG_ID"]

TOKEN_URL = "https://accounts.zoho.com/oauth/v2/token"
API_BASE = "https://www.zohoapis.com/books/v3"

OUT_DIR = Path(__file__).resolve().parent.parent / "output"
OUT_DIR.mkdir(parents=True, exist_ok=True)

COLUMNS = [
    "account_code",
    "account_name",
    "account_type",
    "account_id",
    "parent_account_id",
    "parent_account_name",
    "is_active",
    "debit_total",
    "credit_total",
    "net_debit",
    "net_credit",
    "balance",
    "is_debit",
]


def get_token() -> str:
    r = requests.post(
        TOKEN_URL,
        params={
            "grant_type": "refresh_token",
            "refresh_token": REFRESH_TOKEN,
            "client_id": CLIENT_ID,
            "client_secret": CLIENT_SECRET,
        },
        timeout=30,
    )
    data = r.json()
    if "access_token" not in data:
        raise RuntimeError(f"OAuth failed: {data}")
    return data["access_token"]


def api_get(session: requests.Session, endpoint: str, params: dict[str, Any] | None = None, timeout: int = 180) -> dict:
    params = dict(params or {})
    params["organization_id"] = ORG_ID
    r = session.get(f"{API_BASE}{endpoint}", params=params, timeout=timeout)
    body = r.json()
    if r.status_code != 200 or body.get("code") not in (0, None):
        raise RuntimeError(
            f"{endpoint} failed: HTTP {r.status_code} code={body.get('code')} msg={body.get('message')}"
        )
    return body


def fetch_chart_of_accounts(session: requests.Session) -> dict[str, dict]:
    body = api_get(session, "/chartofaccounts", {"showbalance": "true"})
    rows = body.get("chartofaccounts") or []
    return {str(a["account_id"]): a for a in rows if a.get("account_id")}


def fetch_general_ledger(session: requests.Session, from_date: str, to_date: str) -> list[dict]:
    body = api_get(
        session,
        "/reports/generalledger",
        {
            "filter_by": "TransactionDate.CustomDate",
            "from_date": from_date,
            "to_date": to_date,
        },
    )
    return list(body.get("generalledger") or [])


def _f(val: Any) -> float:
    try:
        return float(val or 0)
    except (TypeError, ValueError):
        return 0.0


def build_rows(gl_rows: list[dict], coa_by_id: dict[str, dict], include_zero: bool) -> list[dict]:
    out: list[dict] = []
    for gl in gl_rows:
        account_id = str(gl.get("account_id") or "")
        debit = round(_f(gl.get("debit_total")), 2)
        credit = round(_f(gl.get("credit_total")), 2)
        balance = round(_f(gl.get("balance")), 2)
        if not include_zero and debit == 0 and credit == 0 and balance == 0:
            continue

        coa = coa_by_id.get(account_id) or {}
        is_debit = bool(gl.get("is_debit")) if "is_debit" in gl else balance >= 0
        if is_debit:
            net_debit, net_credit = abs(balance), 0.0
        else:
            net_debit, net_credit = 0.0, abs(balance)

        out.append(
            {
                "account_code": gl.get("account_code") or coa.get("account_code") or "",
                "account_name": gl.get("name") or coa.get("account_name") or "",
                "account_type": coa.get("account_type") or "",
                "account_id": account_id,
                "parent_account_id": str(coa.get("parent_account_id") or ""),
                "parent_account_name": coa.get("parent_account_name") or "",
                "is_active": coa.get("is_active", True),
                "debit_total": debit,
                "credit_total": credit,
                "net_debit": round(net_debit, 2),
                "net_credit": round(net_credit, 2),
                "balance": balance,
                "is_debit": is_debit,
            }
        )

    out.sort(key=lambda r: (r["account_code"] or "zzzz", r["account_name"]))
    return out


def write_csv(rows: list[dict], path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({c: row.get(c, "") for c in COLUMNS})


def write_excel(rows: list[dict], summary: dict, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "TrialBalance"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row.get(c, "") for c in COLUMNS])
    for i, _ in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = 18

    ws_sum = wb.create_sheet("Summary")
    ws_sum.append(["metric", "value"])
    for cell in ws_sum[1]:
        cell.font = Font(bold=True)
    for k, v in summary.items():
        ws_sum.append([k, v])
    ws_sum.column_dimensions["A"].width = 24
    ws_sum.column_dimensions["B"].width = 40
    wb.save(path)


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Export Zoho Books trial balance (GL + COA)")
    p.add_argument("--from-date", default="2026-04-01", help="YYYY-MM-DD")
    p.add_argument("--to-date", default=date.today().isoformat(), help="YYYY-MM-DD")
    p.add_argument(
        "--include-zero",
        action="store_true",
        help="Keep accounts with zero debit, credit, and balance",
    )
    return p.parse_args()


def main() -> None:
    args = parse_args()
    print(f"Org {ORG_ID}")
    print(f"Period {args.from_date} → {args.to_date}")
    print("Source: /reports/generalledger + /chartofaccounts")

    token = get_token()
    session = requests.Session()
    session.headers.update({"Authorization": f"Zoho-oauthtoken {token}"})

    print("Fetching chart of accounts…")
    coa = fetch_chart_of_accounts(session)
    print(f"  COA accounts: {len(coa)}")

    print("Fetching general ledger…")
    gl = fetch_general_ledger(session, args.from_date, args.to_date)
    print(f"  GL rows: {len(gl)}")

    rows = build_rows(gl, coa, include_zero=args.include_zero)
    summary = {
        "organization_id": ORG_ID,
        "from_date": args.from_date,
        "to_date": args.to_date,
        "source": "GET /reports/generalledger + GET /chartofaccounts",
        "account_count": len(rows),
        "debit_total_sum": round(sum(r["debit_total"] for r in rows), 2),
        "credit_total_sum": round(sum(r["credit_total"] for r in rows), 2),
        "net_debit_sum": round(sum(r["net_debit"] for r in rows), 2),
        "net_credit_sum": round(sum(r["net_credit"] for r in rows), 2),
    }

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    stem = f"trial_balance_{args.from_date}_{args.to_date}_{stamp}"
    csv_path = OUT_DIR / f"{stem}.csv"
    xlsx_path = OUT_DIR / f"{stem}.xlsx"
    json_path = OUT_DIR / f"{stem}.json"

    write_csv(rows, csv_path)
    write_excel(rows, summary, xlsx_path)
    json_path.write_text(json.dumps({"summary": summary, "rows": rows}, indent=2), encoding="utf-8")

    print()
    print(f"Accounts:     {summary['account_count']}")
    print(f"Gross debit:  {summary['debit_total_sum']:,.2f}")
    print(f"Gross credit: {summary['credit_total_sum']:,.2f}")
    print(f"Net debit:    {summary['net_debit_sum']:,.2f}")
    print(f"Net credit:   {summary['net_credit_sum']:,.2f}")
    print(f"Gross balanced: {abs(summary['debit_total_sum'] - summary['credit_total_sum']) < 0.015}")
    print(f"CSV   -> {csv_path}")
    print(f"Excel -> {xlsx_path}")
    print(f"JSON  -> {json_path}")


if __name__ == "__main__":
    main()
